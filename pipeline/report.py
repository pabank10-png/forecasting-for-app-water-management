"""Excel report pack of one run, zipped: a summary workbook plus one workbook per model.

Built by the finish step from the same results HQ shows, then uploaded to the private Supabase Storage bucket
forecast-reports. The website hands admins a short-lived download link (/api/forecast-report).

    Forecast_Report_<date>_<kind>/
      00_สรุปผล.xlsx                ภาพรวม · คะแนนโมเดล · ผลพยากรณ์ · พยากรณ์ทุกโมเดล · น้ำสิ้นเดือน
      task1_reservoir_monthly/<model>.xlsx   วิธีการ · พยากรณ์ · คะแนน · ทดสอบย้อนหลัง · ข้อมูลเข้า
      task1b_reservoir_daily/<model>.xlsx
      README.txt
"""
import datetime as dt
import io
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import data_prep, plan

ROOT = Path(__file__).resolve().parent.parent
TASK_NAME = {plan.T1: "โจทย์ 1: Inflow อ่างรายเดือน", plan.T1B: "โจทย์ 1B: Inflow อ่างรายวัน 1–7 วัน"}
SITE_NAME = {"THREE": "3 อ่างรวม (ดอกกราย+คลองใหญ่+หนองปลาไหล)", "PS": "อ่างประแสร์", "DK": "อ่างดอกกราย",
             "KY": "อ่างคลองใหญ่", "NPL": "อ่างหนองปลาไหล"}
UNIT = {plan.T1: "ล้าน ลบ.ม./เดือน", plan.T1B: "ล้าน ลบ.ม./วัน"}
HEAD = PatternFill("solid", fgColor="1F3864")
WIN = PatternFill("solid", fgColor="E2EFDA")
FONT = "Tahoma"


def _sheet(wb, title, df, note=None, widths=None):
    ws = wb.create_sheet(title[:31])
    row = 1
    if note:
        for line in note if isinstance(note, list) else [note]:
            ws.cell(row, 1, line).font = Font(name=FONT, bold=row == 1, size=12 if row == 1 else 10)
            row += 1
        row += 1
    if df is not None and len(df.columns):
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(row, j, str(c))
            cell.font, cell.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, rec in enumerate(df.itertuples(index=False), row + 1):
            for j, v in enumerate(rec, 1):
                ws.cell(i, j, None if pd.isna(v) else v).font = Font(name=FONT)
        ws.freeze_panes = ws.cell(row + 1, 1)
        for j, c in enumerate(df.columns, 1):
            w = (widths or {}).get(c) or min(40, max(10, len(str(c)) + 2))
            ws.column_dimensions[get_column_letter(j)].width = w
    return ws, row


def _chart(ws, title, header_row, n, cols, anchor, y_title):
    """Line chart over rows header_row+1..header_row+n; cols = column numbers of the series, column 1 = x."""
    if n < 2:
        return
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.height, ch.width = title, y_title, 9, 26
    ch.x_axis.delete = ch.y_axis.delete = False
    for c in cols:
        ch.add_data(Reference(ws, min_col=c, min_row=header_row, max_row=header_row + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + n))
    ws.add_chart(ch, anchor)


def _r(v, d=3):
    return None if v is None or (isinstance(v, float) and pd.isna(v)) else round(float(v), d)


def _readme(name):
    if name.endswith(plan.SUM):
        return ["แยก 3 อ่างแล้วบวก: รันโมเดลนี้กับดอกกราย คลองใหญ่ หนองปลาไหลแยกกัน แล้วบวกเป็น 3 อ่างรวม",
                "P50 = ผลรวม P50 · P10/P90 = ค่ากลางระหว่างแบบ 'ขึ้นลงพร้อมกัน' กับแบบ 'อิสระกัน'",
                "ผลแยกรายอ่างอยู่ในไฟล์ของโมเดลตัวเดียวกัน (ไม่มี +sum)", ""] + _readme(plan.base(name))
    p = ROOT / plan.folder(name) / "README.md"
    lines = p.read_text(encoding="utf-8").splitlines() if p.exists() else [name.split("/")[1]]
    if name in plan.VARIANTS:
        lines = [f"รุ่น inflow อย่างเดียว (ไม่ใส่ฝน ความชื้นดิน ENSO IOD) ของ {plan.folder(name).split('/')[1]}", ""] + lines
    return lines


def _bt_frame(bt):
    if not bt or not bt.get("key"):
        return pd.DataFrame()
    origin, step = zip(*(k.split("|") for k in bt["key"]))
    df = pd.DataFrame({"ต้นทาง (origin)": origin, "ล่วงหน้า (step)": [int(s) for s in step], "ค่าจริง": bt["actual"],
                       "P10": bt["p10"], "P50": bt["p50"], "P90": bt["p90"]})
    df["Error (P50−จริง)"] = df["P50"] - df["ค่าจริง"]
    df["อยู่ในช่วง P10–P90"] = (df["ค่าจริง"] >= df["P10"]) & (df["ค่าจริง"] <= df["P90"])
    return df


def _inputs(task, site):
    try:
        df = data_prep.monthly(site) if task == plan.T1 else data_prep.daily(site)
    except Exception:  # noqa: BLE001 - the report must not fail on a missing input file
        return None
    if task == plan.T1B:
        df = df.tail(3 * 365)  # last 3 years keep the workbook small
    df = df.copy()
    df["date"] = df["date"].astype(str).str[:10]
    return df.round(4)


def model_book(name, result, summary):
    task = name.split("/")[0]
    unit = UNIT[task]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "วิธีการ", None, _readme(name))
    sites = (result or {}).get("sites") or {}
    shown = tuple(plan.PARTS) if name.endswith(plan.SUM) else plan.summary_sites(task)
    gate = {s: (summary["tasks"].get(task, {}).get(s, {}).get("gate") or {}).get(name) for s in shown}
    choice = summary.get("choice", {}).get(task, {})

    rows = []
    for s in shown:
        g, c = gate.get(s) or {}, choice.get(s) or {}
        rows.append({"อ่าง": SITE_NAME[s], "MAE": _r(g.get("mae")), "Skill เทียบ Baseline": _r(g.get("skill")),
                     "Coverage P10–P90": _r(g.get("coverage")),
                     "ผ่านด่าน": {True: "ผ่าน", False: "ไม่ผ่าน", None: "Baseline"}.get(g.get("pass"), "-") if g else "-",
                     "น้ำหนักใน Blend": _r((c.get("weights") or {}).get(name)),
                     "ผู้ชนะ": "✓" if c.get("champion") == name else "",
                     "หมายเหตุ": (result or {}).get("errors", {}).get(s) or (sites.get(s) or {}).get("note")
                     or ("ใช้ผลรอบก่อน" if (sites.get(s) or {}).get("reused") else "")})
    _sheet(wb, "คะแนน", pd.DataFrame(rows),
           [f"คะแนนทดสอบย้อนหลัง ({unit})", "MAE ยิ่งต่ำยิ่งดี · Coverage ควรใกล้ 0.80 · ผ่านด่าน = MAE ต่ำกว่า Baseline และ Coverage อยู่ในช่วง 0.80 ± 2·√(0.16/n) (ราว 0.73–0.87 เมื่อทดสอบ 120 ครั้ง)",
            "รอบรายวันใช้คะแนนจากรอบรายเดือนล่าสุด"], {"หมายเหตุ": 40, "อ่าง": 36})

    for s, r in sites.items():
        fc = pd.DataFrame(r.get("forecast") or [])
        if len(fc):
            fc = fc.rename(columns={"step": "ล่วงหน้า", "date": "วันที่"})[["ล่วงหน้า", "วันที่", "p10", "p50", "p90"]]
            ws, hr = _sheet(wb, f"พยากรณ์_{s}", fc, [f"{SITE_NAME[s]}: ผลพยากรณ์ ({unit})"])
            _chart(ws, "P10 / P50 / P90", hr, len(fc), [3, 4, 5], "H3", unit)
            for src, rows_ in (r.get("by_source") or {}).items():
                fs = pd.DataFrame(rows_).rename(columns={"step": "ล่วงหน้า", "date": "วันที่"})
                _sheet(wb, f"ฝน{src}_{s}", fs, [f"{SITE_NAME[s]}: ใช้ฝนพยากรณ์ {src}"])
        bt = _bt_frame(r.get("bt"))
        if len(bt):
            ws, hr = _sheet(wb, f"ย้อนหลัง_{s}", bt, [f"{SITE_NAME[s]}: ทดสอบย้อนหลัง ({unit})",
                                                     "แต่ละแถว = ณ วันต้นทาง พยากรณ์ล่วงหน้า step แล้วเทียบค่าจริง"])
            one = bt[bt["ล่วงหน้า (step)"] == 1].reset_index(drop=True)
            if len(one):
                c0 = len(bt.columns) + 2
                for j, col in enumerate(["ต้นทาง (origin)", "ค่าจริง", "P10", "P50", "P90"]):
                    ws.cell(hr, c0 + j, col).font = Font(name=FONT, bold=True)
                    for i, v in enumerate(one[col], hr + 1):
                        ws.cell(i, c0 + j, None if pd.isna(v) else v)
                ch = LineChart()
                ch.title, ch.height, ch.width = "ล่วงหน้า 1 ช่วง: ค่าจริง เทียบ P10/P50/P90", 9, 28
                ch.x_axis.delete = ch.y_axis.delete = False
                for j in range(1, 5):
                    ch.add_data(Reference(ws, min_col=c0 + j, min_row=hr, max_row=hr + len(one)), titles_from_data=True)
                ch.set_categories(Reference(ws, min_col=c0, min_row=hr + 1, max_row=hr + len(one)))
                ws.add_chart(ch, f"{get_column_letter(c0 + 6)}3")
    for s in shown:
        df = _inputs(task, s)
        if df is not None and (s in sites or not sites):
            _sheet(wb, f"ข้อมูลเข้า_{s}", df, [f"{SITE_NAME[s]}: ข้อมูลที่โมเดลอ่าน" +
                                              (" (3 ปีล่าสุด)" if task == plan.T1B else "")])
    if not sites:
        _sheet(wb, "ไม่มีผล", None, ["โมเดลนี้ไม่มีผลในรอบนี้", str((result or {}).get("errors") or "")])
    return wb

METHOD = [  # (หัวข้อ, อธิบาย / สูตร, ตัวอย่าง)
    ("ภาพรวม", "ทุกโมเดลทายเป็นช่วง P10 / P50 / P90 แล้วถูกทดสอบย้อนหลัง (ทายจากข้อมูลเท่าที่มีในตอนนั้น เทียบค่าจริง) "
     "คะแนนที่ได้ใช้ตัดสิน 3 ขั้น: 1) ผ่านด่านไหม 2) ได้น้ำหนักเท่าไร 3) ใช้ตัวเดียวหรือ Blend",
     "รอบรายเดือนคิดคะแนนใหม่ทุกครั้ง รอบรายวันใช้ตัวที่รอบรายเดือนล่าสุดเลือกไว้ จึงไม่ใช่การเลือกโมเดลแบบตายตัว"),
    ("1. P10 / P50 / P90", "Percentile ของค่าที่เป็นไปได้\n"
     "P10 = มีโอกาส 10% ที่ค่าจริงต่ำกว่านี้ (กรณีน้ำน้อย)\nP50 = ค่ากลาง โอกาสต่ำกว่า/สูงกว่าเท่ากัน ใช้คิด MAE\n"
     "P90 = มีโอกาส 10% ที่ค่าจริงสูงกว่านี้ (กรณีน้ำมาก)\nช่วง P10–P90 จึงควรครอบค่าจริง 100 − 10 − 10 = 80% ของครั้ง",
     "ทาย P10 70, P50 106, P90 184 ล้าน ลบ.ม. = ค่ากลาง 106 และ 8 ใน 10 ครั้งค่าจริงควรอยู่ระหว่าง 70–184"),
    ("2. วิธีหา P: Climatology", "เอาค่าเดือนเดียวกันทุกปีมาเรียงน้อยไปมาก แล้วอ่านค่าที่ตำแหน่ง 10% / 50% / 90% ตรงๆ (ไม่ใช้สูตร Normal)\n"
     "ตำแหน่ง = (n − 1) × p นับตัวแรกเป็น 0 ถ้าตกระหว่างสองตัวให้เฉลี่ยตามสัดส่วน",
     "inflow ต.ค. 3 อ่างรวม 20 ปี (2006–2025) เรียงแล้ว: 43.2, 63.3, 73.2, ... , 104.6, 107.4, ... , 183.1, 190.3, 316.5\n"
     "P10: 19×0.1 = 1.9 → 63.3 + 0.9×(73.2−63.3) = 72.2\nP50: 19×0.5 = 9.5 → (104.6+107.4)/2 = 106.0\n"
     "P90: 19×0.9 = 17.1 → 183.1 + 0.1×(190.3−183.1) = 183.8"),
    ("2. วิธีหา P: Seasonal naive", "P50 = ค่าเดือนเดียวกันของปีก่อน\nP10/P90 = P50 + Percentile 10/90 ของ Error ในอดีต (ค่าจริง − ค่าปีก่อน)", ""),
    ("2. วิธีหา P: SARIMAX / ETS (สูตร Normal)", "Pₚ = μ + zₚ × σ   (μ = ค่าเฉลี่ย, σ = SD, zₚ = Φ⁻¹(p) ค่าคงที่ตาม Percentile)\n"
     "z: P10 = −1.2816, P50 = 0, P90 = +1.2816 (P5/P95 = ±1.645, P2.5/P97.5 = ±1.96)\n"
     "Excel: =NORM.INV(0.9, ค่าเฉลี่ย, SD) หรือ z = NORM.S.INV(0.9)\n"
     "inflow เบ้ (ติดลบไม่ได้ แต่บางปีพุ่งสูง) จึงคิดใน log ก่อน: Pₚ = exp(μ_log + zₚ × σ_log) − 1",
     "Normal ตรงๆ: μ 20, σ 8 → P10 = 20 − 1.2816×8 = 9.75, P90 = 30.25\n"
     "แบบ log: μ_log 3.0, σ_log 0.5 → P10 9.6, P50 19.1, P90 37.1 (ยืดไปทางน้ำมาก)\n"
     "ข้อมูล ต.ค. ข้างบนถ้าใช้ Normal ตรงๆ (μ 119.5, σ 62.2) ได้ P10 39.8 / P50 119.5 / P90 199.2 ผิดจากจริงมาก เพราะข้อมูลเบ้"),
    ("2. วิธีหา P: LightGBM / XGBoost / N-BEATS / N-HiTS / LSTM", "ฝึกให้ทายแต่ละ P โดยตรงด้วย Pinball loss\n"
     "Loss = max(q × e, (q − 1) × e)   e = ค่าจริง − ค่าทาย, q = 0.1 / 0.5 / 0.9",
     "P90 (q = 0.9): ทายต่ำไป 10 โดนปรับ 9 แต่ทายสูงไป 10 โดนแค่ 1 โมเดลจึงเรียนรู้ที่จะทายสูงไว้ จนค่าจริงเกินแค่ราว 10% ของครั้ง"),
    ("2. วิธีหา P: Chronos-2 / TimesFM / Moirai", "Foundation model ให้ผลเป็น Percentile ออกมาเองอยู่แล้ว หยิบระดับ 0.1 / 0.5 / 0.9 มาใช้", ""),
    ("2. วิธีหา P: Conformal", "ปรับช่วงของโมเดลฐานให้ครอบค่าจริงพอดี 80%\ns = max(P10 − ค่าจริง, ค่าจริง − P90) ของทุกครั้งในอดีต\n"
     "q = ค่า s ที่ทำให้ 80% ของอดีตอยู่ในช่วง → ช่วงใหม่ = [P10 − q, P90 + q]", ""),
    ("3. MAE", "MAE = ค่าเฉลี่ยของ |ค่าจริง − P50| ทุกครั้งที่ทดสอบ (ช่วงล่วงหน้า 1–3 เดือน สำหรับโจทย์ 1) ยิ่งต่ำยิ่งแม่น หน่วยเดียวกับข้อมูล",
     "ค่าจริง 100, 50 ทาย 90, 60 → MAE = (10 + 10)/2 = 10"),
    ("4. Skill", "Skill = 1 − MAE โมเดล ÷ MAE Baseline (Baseline = Climatology)\n0 = เท่า Baseline, บวก = ดีกว่า, ลบ = แย่กว่าค่าเฉลี่ยในอดีต",
     "Chronos-2 3 อ่างรวม MAE 15.05 เทียบ Climatology 16.17 → Skill = 1 − 15.05/16.17 = 0.07 (แม่นกว่า 7%)"),
    ("5. Coverage", "Coverage = สัดส่วนครั้งที่ค่าจริงอยู่ในช่วง P10–P90 ควรใกล้ 0.80 (ตรงตามนิยามในข้อ 1)\n"
     "ต่ำกว่า = ช่วงแคบ มั่นใจเกินจริง, สูงกว่า = ช่วงกว้างจนใช้ตัดสินใจไม่ได้\n"
     "หลักสากล: Coverage ที่วัดได้ต้องใกล้ระดับที่ประกาศ งานพยากรณ์น้ำนิยมช่วง 10/50/90 (เช่น NWS สหรัฐฯ, BoM ออสเตรเลีย)",
     "ทดสอบ 120 ครั้ง ค่าจริงอยู่ในช่วง 96 ครั้ง → Coverage = 96/120 = 0.80"),
    ("6. ผ่านด่าน", "ต้องผ่านทั้ง 2 ข้อ:\n1) MAE ต่ำกว่า Baseline\n2) Coverage อยู่ในช่วง 0.80 ± 2 × √(0.8 × 0.2 ÷ n)  "
     "(n = จำนวนจุดเริ่มทดสอบ, 2 SD ของความคลาดเคลื่อนตามโอกาสแบบ binomial = ช่วงที่โมเดลแม่นจริงจะตกใน 95% ของกรณี)\n"
     "ไม่มีโมเดลผ่าน → ใช้ Baseline",
     "โจทย์ 1: ทดสอบทุกเดือน 2016–2025 = 120 ครั้ง → SD = √(0.16/120) = 0.037 → ช่วง 0.727–0.873\n"
     "Coverage 0.85 ผ่าน, 0.95 ไม่ผ่าน (กว้างเกิน), 0.60 ไม่ผ่าน (แคบเกิน)\nโจทย์ 1B ทดสอบหลายครั้งกว่า ช่วงจึงแคบลงเอง"),
    ("7. น้ำหนัก (Blend)", "เฉพาะโมเดลที่ผ่านด่าน\nน้ำหนักดิบ wᵢ = (1/MAEᵢ) ÷ Σ(1/MAEⱼ)\n"
     "ดึงเข้าหาค่าเท่ากันครึ่งหนึ่ง (Shrinkage, λ = 0.5): wᵢ' = 0.5 × wᵢ + 0.5 × (1/จำนวนโมเดล) กันไม่ให้เชื่อตัวเดียวมากเกิน\n"
     "Blend = Σ wᵢ' × ค่าพยากรณ์ของโมเดล i",
     "MAE A 12, B 15, C 18 → 1/MAE = 0.0833, 0.0667, 0.0556 (รวม 0.2056)\nน้ำหนักดิบ 0.405 / 0.324 / 0.270\n"
     "หลังดึงเข้าหา 0.333: 0.369 / 0.329 / 0.302\nเดือนหน้า A 100, B 110, C 130 → Blend = 0.369×100 + 0.329×110 + 0.302×130 = 112.4"),
    ("8. กฎ 2% (เลือกตัวเดียวหรือ Blend)", "ทดสอบ Blend ย้อนหลังแบบเดียวกัน ใช้ Blend ก็ต่อเมื่อ MAE Blend ≤ MAE ตัวเดี่ยวที่ดีที่สุด × 0.98\n"
     "ไม่ถึง → ใช้แชมป์ตัวเดียว (น้ำหนัก 1.0) เพราะดีขึ้นนิดเดียวไม่คุ้มความซับซ้อน",
     "ตัวเดี่ยวดีสุด A MAE 12 → เกณฑ์ 11.76\nBlend 11.5 → ใช้ Blend (112.4)\nBlend 11.9 → ใช้ A ตัวเดียว (100)"),
    ("9. แบบ +sum (3 อ่างรวม)", "นอกจากทาย 3 อ่างรวมตรงๆ ยังทายดอกกราย คลองใหญ่ หนองปลาไหล แยกกันแล้วบวก ชื่อโมเดลลงท้าย +sum และแข่งในด่านเดียวกัน\n"
     "P50 = ผลรวม P50, P10/P90 = ค่ากลางระหว่างการบวกแบบไปทางเดียวกันทั้งหมด กับแบบอิสระต่อกัน", ""),
]


def method_sheet(wb, title="วิธีคิด", index=None):
    """Plain-language sheet: what P10/P50/P90, MAE, Skill, Coverage, the gate and the blend weights mean."""
    ws, hr = _sheet(wb, title, pd.DataFrame(METHOD, columns=["หัวข้อ", "อธิบาย / สูตร", "ตัวอย่าง"]),
                    ["วิธีคิด: P10/P50/P90, MAE, Skill, Coverage, ผ่านด่าน และน้ำหนัก",
                     "อ่านจากบนลงล่าง ตัวอย่างตัวเลขบางข้อเป็นตัวเลขสมมติเพื่ออธิบาย"],
                    {"หัวข้อ": 30, "อธิบาย / สูตร": 80, "ตัวอย่าง": 70})
    for row in ws.iter_rows(min_row=hr + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if index is not None:
        wb.move_sheet(ws, offset=index - wb.index(ws))
    return ws


def summary_book(summary, run):
    wb = Workbook()
    wb.remove(wb.active)
    info = [f"รายงานผลพยากรณ์ Water Intelligence · รอบ{'รายเดือน' if summary.get('kind') == 'monthly' else 'รายวัน'}",
            f"เสร็จเมื่อ {run.get('finished')} (เวลาไทย) · run id {run.get('id')}",
            f"GitHub Actions: {run.get('url') or '-'}"]
    over, scores, finals, every = [], [], [], []
    for task, sites in summary.get("tasks", {}).items():
        for s, e in sites.items():
            c = summary.get("choice", {}).get(task, {}).get(s) or {}
            base = (e.get("gate") or {}).get(plan.BASELINE[task]) or {}
            champ = c.get("champion")
            cg = (e.get("gate") or {}).get(champ) or {}
            over.append({"โจทย์": TASK_NAME[task], "อ่าง": SITE_NAME[s],
                         "ผู้ชนะ": "Blend (Shrinkage)" if champ == "ensemble" else (champ or "-").split("/")[-1],
                         "MAE ผู้ชนะ": _r(c.get("blend_mae") if champ == "ensemble" else cg.get("mae")),
                         "MAE Baseline": _r(base.get("mae")), "ผ่านด่าน (ตัว)": len(c.get("members") or []),
                         "น้ำหนัก": ", ".join(f"{k.split('/')[-1]} {v:.2f}" for k, v in (c.get("weights") or {}).items()),
                         "หน่วย": UNIT[task]})
            for name, g in sorted((e.get("gate") or {}).items(), key=lambda kv: kv[1].get("mae") or 1e9):
                scores.append({"โจทย์": task, "อ่าง": s, "โมเดล": name.split("/")[-1], "MAE": _r(g.get("mae")),
                               "Skill": _r(g.get("skill")), "Coverage": _r(g.get("coverage")),
                               "ผ่านด่าน": {True: "ผ่าน", False: "ไม่ผ่าน", None: "Baseline"}.get(g.get("pass")),
                               "น้ำหนัก": _r((c.get("weights") or {}).get(name)),
                               "ผู้ชนะ": "✓" if champ == name else ""})
            for r in e.get("final") or []:
                finals.append({"โจทย์": task, "อ่าง": s, "ล่วงหน้า": r["step"], "วันที่": r["date"],
                               "P10": r["p10"], "P50": r["p50"], "P90": r["p90"]})
            for name, m in (e.get("models") or {}).items():
                for r in m.get("forecast") or []:
                    every.append({"โจทย์": task, "อ่าง": s, "โมเดล": name.split("/")[-1], "ล่วงหน้า": r["step"],
                                  "วันที่": r["date"], "P10": r["p10"], "P50": r["p50"], "P90": r["p90"]})
    _sheet(wb, "ภาพรวม", pd.DataFrame(over), info, {"โจทย์": 30, "อ่าง": 36, "น้ำหนัก": 40})
    method_sheet(wb)
    _sheet(wb, "คะแนนโมเดล", pd.DataFrame(scores),
           ["คะแนนทดสอบย้อนหลังทุกโมเดล เรียงจาก MAE ต่ำสุด", "ผ่านด่าน = MAE ต่ำกว่า Baseline และ Coverage อยู่ในช่วง 0.80 ± 2·√(0.16/n) (ราว 0.73–0.87 เมื่อทดสอบ 120 ครั้ง)"])
    ws, hr = _sheet(wb, "ผลพยากรณ์สุดท้าย", pd.DataFrame(finals), ["ผลที่เว็บใช้ (ผู้ชนะของแต่ละอ่าง)"])
    _sheet(wb, "พยากรณ์ทุกโมเดล", pd.DataFrame(every), ["ผลพยากรณ์ของทุกโมเดล"])
    st = [{"อ่าง": SITE_NAME[s], "วันที่": e["storage"].get("date"), "น้ำในอ่างวันนี้": e["storage"].get("storage_today"),
           "Inflow สะสมเดือนนี้": e["storage"].get("inflow_month_to_date"), "ความจุ": e["storage"].get("capacity"),
           "Hard Red": e["storage"].get("dead"), "MAE inflow 1–3 เดือน": _r(e["storage"].get("inflow_mae"))}
          for s, e in summary.get("tasks", {}).get(plan.T1, {}).items() if e.get("storage")]
    _sheet(wb, "น้ำสิ้นเดือน_จุดเริ่ม", pd.DataFrame(st),
           ["จุดเริ่มคำนวณน้ำสิ้นเดือน (ล้าน ลบ.ม.)", "Outflow และน้ำผันกรอกเองในหน้า HQ ส่วน Publish"], {"อ่าง": 36})
    return wb


README = """รายงานผลพยากรณ์ Water Intelligence

00_สรุปผล.xlsx
  ภาพรวม              ผู้ชนะของแต่ละอ่าง, MAE เทียบ Baseline, น้ำหนัก Blend
  วิธีคิด              P10/P50/P90, MAE, Skill, Coverage, ผ่านด่าน, น้ำหนัก อธิบายพร้อมตัวอย่าง
  คะแนนโมเดล          ทุกโมเดล ทุกอ่าง เรียงจาก MAE ต่ำสุด
  ผลพยากรณ์สุดท้าย     ค่าที่เว็บแสดง (P10/P50/P90)
  พยากรณ์ทุกโมเดล      ผลของทุกโมเดลไว้เทียบกัน
  น้ำสิ้นเดือน_จุดเริ่ม   ปริมาณน้ำวันนี้และ inflow สะสม ใช้คำนวณน้ำสิ้นเดือนในหน้า HQ

<โจทย์>/<โมเดล>.xlsx
  วิธีการ        หลักการ สูตร ตัวแปร (จาก README ของโมเดล)
  คะแนน          MAE, Skill, Coverage, ผ่านด่าน, น้ำหนัก ของแต่ละอ่าง
  พยากรณ์_<อ่าง>  P10/P50/P90 พร้อมกราฟ (โจทย์ 1B มีแยกตามค่ายฝนพยากรณ์ด้วย)
  ย้อนหลัง_<อ่าง>  ผลทดสอบย้อนหลังทุกจุด พร้อมกราฟค่าจริงเทียบพยากรณ์ (เฉพาะรอบรายเดือน)
  ข้อมูลเข้า_<อ่าง> ข้อมูลที่โมเดลอ่านในรอบนี้

รอบรายวันไม่มีการทดสอบย้อนหลัง ใช้คะแนนจากรอบรายเดือนล่าสุด
"""


def build(results, summary, run):
    """Zip bytes and the file name. run: {id, kind, finished, url}"""
    day = run["finished"][:10]
    folder = f"Forecast_Report_{day}_{summary.get('kind')}"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        def put(path, wb):
            b = io.BytesIO()
            wb.save(b)
            z.writestr(f"{folder}/{path}", b.getvalue())
        put("00_สรุปผล.xlsx", summary_book(summary, run))
        for name in plan.MODELS:
            if name in results or any(name in (e.get("models") or {}) for e in summary["tasks"].get(name.split("/")[0], {}).values()):
                r = results.get(name) or {"sites": {s: e["models"][name] for s, e in summary["tasks"][name.split("/")[0]].items()
                                                    if name in (e.get("models") or {})}}
                put(f"{name}.xlsx", model_book(name, r, summary))
        z.writestr(f"{folder}/README.txt", README)
    return buf.getvalue(), f"{day}_{summary.get('kind')}_{run['id'][:8]}.zip"


def thai_now():
    return (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=7)).strftime("%Y-%m-%d %H:%M")
